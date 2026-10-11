"""Import the credit team's reviewed golden questions (evals/golden_review.xlsx, sheet "Questions") into the golden
tests. Only rows marked Approved = Yes are imported; a question already in the golden tests is updated.

    python -m evals.import_golden_review                       # evals/golden_review.xlsx
    python -m evals.import_golden_review path/to/reviewed.xlsx
    python -m evals.import_golden_review --dry                 # check the file, change nothing
"""
import sys
from pathlib import Path
from typing import List, Tuple

DEFAULT = Path(__file__).with_name("golden_review.xlsx")
SHEET = "Questions"
NOT_FOUND = "(not found)"
# workbook header -> golden-test field
COLUMNS = {"ID": "review_id", "Category": "category", "Question": "question",
           "Earlier questions (follow-ups)": "history", "Loan type": "loan_type",
           "Should the app answer?": "should_answer", "Reference answer": "reference_answer",
           "Must contain": "expected_answer", "Expected document": "expected_document", "Expected page": "expected_page",
           "Reviewed by": "approved_by", "Approved?": "approved", "Reviewer notes": "review_notes"}


def _text(v) -> str:
    return "" if v is None else str(v).strip()


def read(path: Path) -> Tuple[List[dict], List[str]]:
    """Approved rows as golden-test items, and the problems found (rows with problems are not imported)."""
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)[SHEET]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {COLUMNS[_text(h)]: i for i, h in enumerate(header) if _text(h) in COLUMNS}
    missing = set(COLUMNS.values()) - set(idx)
    if missing:
        raise ValueError(f"Sheet '{SHEET}' is missing columns: {', '.join(sorted(missing))}")
    items, problems = [], []
    for n, row in enumerate(rows, 2):
        r = {k: _text(row[i]) if i < len(row) else "" for k, i in idx.items()}
        if not r["question"] or r["approved"].lower() not in ("yes", "y"):
            continue
        where = f"row {n} ({r['review_id'] or r['question'][:40]})"
        nothing = r["should_answer"].lower().startswith("not found")
        if not nothing and not r["reference_answer"]:
            problems.append(f"{where}: approved without a reference answer")
            continue
        if not r["approved_by"]:
            problems.append(f"{where}: approved without 'Reviewed by'")
            continue
        page = r["expected_page"]
        if page and not page.split(".")[0].isdigit():
            problems.append(f"{where}: expected page '{page}' is not a number")
            continue
        items.append({
            "question": r["question"], "category": r["category"], "history": r["history"],
            "loan_type": r["loan_type"], "approved_by": r["approved_by"],
            "reference_answer": NOT_FOUND if nothing else r["reference_answer"],
            "expected_answer": NOT_FOUND if nothing else r["expected_answer"],
            "expected_document": "" if nothing else r["expected_document"],
            "expected_page": int(page.split(".")[0]) if page else None,
            "notes": "; ".join(x for x in (f"review: {r['review_id']}", r["review_notes"]) if x),
            "active": 1})
    return items, problems


def apply(items: List[dict], dry: bool = False) -> Tuple[int, int]:
    from app.services import rag_debug
    rag_debug.init_db()
    have = {" ".join(g["question"].lower().split()): g["id"] for g in rag_debug.golden_list()}
    added = updated = 0
    for it in items:
        gid = have.get(" ".join(it["question"].lower().split()))
        if not dry:
            rag_debug.golden_save(it, gid)
        added, updated = (added, updated + 1) if gid else (added + 1, updated)
    return added, updated


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry" in sys.argv
    path = Path(args[0]) if args else DEFAULT
    items, problems = read(path)
    for p in problems:
        print("skipped  " + p)
    added, updated = apply(items, dry)
    print(f"\n{len(items)} approved row(s): {added} {'to add' if dry else 'added'}, {updated} "
          f"{'to update' if dry else 'updated'}; {len(problems)} skipped.")


if __name__ == "__main__":
    main()
