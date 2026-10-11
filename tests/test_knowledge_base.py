"""Urban Money folder: multi-format extraction, incremental sync (add / change / delete) and hybrid search."""
import json
from pathlib import Path

import pytest
from langchain_core.documents import Document

import app.rag.vectorstore as vectorstore
from app.core.config import get_settings
from app.services import knowledge_base as kb
from app.services.ingestion import extract, load_and_chunk


class FakeStore:
    """Records vector writes; stands in for Pinecone (persistent) without embeddings or network."""

    def __init__(self):
        self.docs = {}

    def add_documents(self, docs, ids=None):
        self.docs.update(zip(ids, docs))

    def delete(self, ids=None):
        for i in ids:
            self.docs.pop(i, None)


@pytest.fixture
def folder(monkeypatch, tmp_path):
    root = tmp_path / "Urban Money"
    (root / "sub").mkdir(parents=True)
    store = FakeStore()
    monkeypatch.setattr(get_settings(), "urban_money_dir", root)
    monkeypatch.setattr(get_settings(), "kb_index_db", tmp_path / "kb.db")
    monkeypatch.setattr(get_settings(), "kb_cache_dir", tmp_path / "cache")
    monkeypatch.setattr(vectorstore, "is_pinecone", lambda: True)
    monkeypatch.setattr(vectorstore, "get_vectorstore", lambda: store)
    return root, store


def _xlsx(path: Path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Codes"
    for r in rows:
        ws.append(r)
    wb.save(path)


def test_spreadsheet_rows_keep_headers_and_title(tmp_path):
    f = tmp_path / "Code Book.xlsx"
    _xlsx(f, [["CODE BOOK URBAN MONEY"], ["BANK NAME", "HL", "LAP"], ["Axis Bank", "6159MUM", "6159MUM"],
              ["Bajaj Finserv", None, "14788"]])
    docs = load_and_chunk(f)
    text = docs[0].page_content
    assert text.startswith("Source: Code Book.xlsx | Sheet: Codes\nCODE BOOK URBAN MONEY")
    assert "BANK NAME: Axis Bank | HL: 6159MUM | LAP: 6159MUM" in text
    assert "BANK NAME: Bajaj Finserv | LAP: 14788" in text


def test_csv_json_html_vcf_are_extracted(tmp_path):
    (tmp_path / "pins.csv").write_text("area_pincode\n400001\n400002\n")
    (tmp_path / "x.json").write_text(json.dumps([{"lender": "Axis", "code": {"hl": "A1"}}]))
    (tmp_path / "p.html").write_text("<html><script>x()</script><h1>LAP Policy</h1><p>Min CIBIL 700</p></html>")
    (tmp_path / "rm.vcf").write_text("BEGIN:VCARD\nFN:Ravi RM\nTEL;type=CELL:+91 90000 11111\nEND:VCARD\n")
    pins = extract(tmp_path / "pins.csv")[0]
    assert pins.context == "area_pincode" and pins.rows == ["400001", "400002"]
    assert extract(tmp_path / "x.json")[0].rows == ["lender: Axis | code.hl: A1"]
    html = extract(tmp_path / "p.html")[0]
    assert html.heading == "LAP Policy" and "Min CIBIL 700" in html.text and "x()" not in html.text
    assert extract(tmp_path / "rm.vcf")[0].text == "Name: Ravi RM\nPhone: +91 90000 11111"


def test_sync_adds_updates_and_removes_files(folder):
    root, store = folder
    (root / "notes.txt").write_text("Piramal LAP payout is 1.5 percent.")
    _xlsx(root / "sub" / "codes.xlsx", [["BANK NAME", "CODE"], ["Axis Bank", "6159MUM"]])
    (root / "voice.opus").write_bytes(b"\x00")
    r = kb.sync()
    assert sorted(r["indexed"]) == ["notes.txt", "sub/codes.xlsx"] and r["skipped"] == ["voice.opus (unsupported type)"]
    assert all(d.metadata["corpus"] == "urban_money" for d in store.docs.values())
    first_ids = set(store.docs)

    r = kb.sync()                                                # nothing changed
    assert r["indexed"] == [] and r["unchanged"] == 2 and set(store.docs) == first_ids

    (root / "notes.txt").write_text("Piramal LAP payout is 1.75 percent from October.")
    (root / "sub" / "codes.xlsx").unlink()
    r = kb.sync()
    assert r["indexed"] == ["notes.txt"] and r["removed"] == ["sub/codes.xlsx"]
    texts = [d.page_content for d in store.docs.values()]
    assert len(texts) == 1 and "1.75 percent" in texts[0]       # old version and deleted file are gone
    assert kb.keyword_search("code 6159MUM") == []                # keyword index pruned too


def test_unreadable_file_is_reported_not_fatal(folder):
    root, _ = folder
    (root / "poster.png").write_bytes(b"not really an image")    # OCR is disabled in tests
    (root / "ok.md").write_text("Bandhan Bank BL min vintage 3 years")
    r = kb.sync()
    assert r["indexed"] == ["ok.md"] and r["failed"] and r["failed"][0].startswith("poster.png")
    assert kb.sync()["failed"]                                   # still reported, not retried until it changes


def test_keyword_search_finds_exact_codes_and_pincodes(folder):
    root, _ = folder
    (root / "pins.csv").write_text("Pincode,City\n411038,Pune\n560001,Bangalore\n")
    kb.sync()
    hits = kb.keyword_search("is 560001 serviceable?")
    assert hits and "560001" in hits[0].page_content and hits[0].metadata["source"] == "pins.csv"
    assert kb.keyword_search("home loan for salaried") == []    # no exact terms -> vectors only


def test_lexical_search_finds_names_in_large_sheets(folder):
    root, _ = folder
    rows = [["COMPANY NAME", "NAME", "CONTACT NUMBER"]] + [[f"LENDER {i}", f"RM {i}", f"90000{i:05d}"] for i in range(300)]
    _xlsx(root / "bankers.xlsx", rows[:150] + [["CHOLA", "AMIT MISHRA", "8956266986"]] + rows[150:])
    kb.sync()
    hit = kb.lexical_search("Chola personal loan banker contact", k=1)[0].page_content
    assert "COMPANY NAME: CHOLA | NAME: AMIT MISHRA | CONTACT NUMBER: 8956266986" in hit


def test_huge_file_goes_to_keyword_index_only(folder, monkeypatch):
    root, store = folder
    rows = [["Code", "Employer Name", "Category"]] + [[300000 + i, f"EMPLOYER {i} PVT LTD", "PB"] for i in range(400)]
    rows.insert(200, [399999, "TBWA INDIA PRIVATE LIMITED", "CAT A"])
    _xlsx(root / "company list.xlsx", rows)
    (root / "notes.txt").write_text("Piramal LAP payout is 1.5 percent.")
    path = root / "company list.xlsx"
    sha = kb._hash_file(path)
    n = len(kb.chunk_file(path, root, sha))
    stale = kb.chunk_ids("company list.xlsx", sha, n)               # an earlier, interrupted vector upload
    store.add_documents([Document(page_content="old")] * 3, ids=stale[:3])
    with kb.closing(kb._db()) as db:
        db.execute("INSERT INTO kb_files VALUES (?,?,?,?,?,?,?,?,?)", (kb._store_key(), "company list.xlsx", sha,
                   path.stat().st_size, path.stat().st_mtime, n, "indexing", None, kb._now()))
        db.commit()
    monkeypatch.setattr(get_settings(), "kb_vector_max_chunks", n - 1)

    r = kb.sync()
    assert r["keyword_only"] == ["company list.xlsx"] and r["indexed"] == ["notes.txt"]
    assert all("payout" in d.page_content for d in store.docs.values())  # stale vectors gone, no new ones
    hit = kb.lexical_search("is TBWA India on the list", k=1)[0].page_content
    assert "Employer Name: TBWA INDIA PRIVATE LIMITED | Category: CAT A" in hit
    assert kb.keyword_search("code 399999")
    assert kb.files()[0]["status"] == "keyword"

    r = kb.sync()                                                   # unchanged: not redone
    assert r["keyword_only"] == [] and r["unchanged"] == 2
    path.unlink()
    assert kb.sync()["removed"] == ["company list.xlsx"] and kb.keyword_search("code 399999") == []


def test_search_ranks_urban_money_first(monkeypatch):
    um = Document(page_content="Source: x.xlsx\nAxis HL code 6159MUM", metadata={"corpus": "urban_money"})
    other = Document(page_content="Sample Lender A HL policy", metadata={"loan_type": "Home Loan"})

    class VS:
        def similarity_search_with_score(self, q, k, filter=None):       # other source scores higher raw
            return [(d, s) for d, s in ((um, 0.50), (other, 0.55)) if filter(d)]

    monkeypatch.setattr(vectorstore, "get_vectorstore", lambda: VS())
    monkeypatch.setattr(vectorstore, "is_pinecone", lambda: False)
    out = vectorstore.search("home loan code", "Home Loan", k=2)
    assert [d.page_content for d in out] == [um.page_content, other.page_content]


def test_search_keeps_matched_lender_text(monkeypatch):
    ums = [Document(page_content=f"Source: x.xlsx\nrow {i}", metadata={"corpus": "urban_money"}) for i in range(6)]
    own = Document(page_content="Lender: Sample Lender A\nHL policy", metadata={"lender": "Sample Lender A",
                                                                                    "loan_type": "Home Loan"})

    class VS:
        def similarity_search_with_score(self, q, k, filter=None):
            return [(d, 0.9) for d in ums if filter(d)][:k] + [(own, 0.2)] * filter(own)

    monkeypatch.setattr(vectorstore, "get_vectorstore", lambda: VS())
    monkeypatch.setattr(vectorstore, "is_pinecone", lambda: False)
    out = vectorstore.search("Home Loan policy eligibility", "Home Loan", ["Sample Lender A"], k=4)
    assert out[0] is own and len(out) == 4


def test_folder_found_case_insensitively(monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), "urban_money_dir", None)
    monkeypatch.setattr(kb, "BASE_DIR", tmp_path)
    (tmp_path / "data" / "URBAN MONEY").mkdir(parents=True)
    assert kb.kb_root() == tmp_path / "data" / "URBAN MONEY"
