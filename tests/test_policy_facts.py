"""Structured policy facts: imported as pending, used only once approved and in force, ahead of document text."""
from datetime import date

import pytest

import app.rag.chat as chat
import app.rag.retrieval as retrieval
from app.services import policy_facts
from tests.test_kb_admin import _admin

CSV = """lender,product,fact,value,unit,customer_category,effective_from,effective_until,source,page
HDFC LTD,HL,roi,8.50% - 9.40%,,Salaried,2026-08-01,,HDFC HL rate sheet Aug 2026.pdf,2
HDFC Bank,Home Loan,roi,9.10%,,Salaried,2026-01-01,2026-07-31,HDFC HL rate sheet Jan 2026.pdf,2
ICICI Bank,Home Loan,payout,0.65%,,,2026-08-01,,AUG 26 PAYOUT SHEET.xls,
"""


@pytest.fixture(autouse=True)
def clean():
    yield
    policy_facts.delete([f["fact_id"] for f in policy_facts.listing()])


def test_import_validates_then_waits_for_approval():
    with pytest.raises(ValueError, match="Row 2: fact must be one of"):
        policy_facts.import_csv("lender,fact,value\nHDFC,interest,9%\n")
    with pytest.raises(ValueError, match="effective_from must be a date"):
        policy_facts.import_csv("lender,fact,value,effective_from\nHDFC,roi,9%,Aug 26\n")
    assert policy_facts.import_csv(CSV, "admin@example.com") == {"rows": 3, "added": 3, "skipped": 0}
    assert policy_facts.import_csv(CSV)["skipped"] == 3                          # same rows again
    assert policy_facts.current(["HDFC Bank"], ["roi"]) == []                   # pending: not used
    policy_facts.approve([f["fact_id"] for f in policy_facts.listing("pending")])
    rows = policy_facts.current(["HDFC Bank"], ["roi"], "Home Loan", today=date(2026, 10, 9))
    assert [(r["lender"], r["product"], r["value"]) for r in rows] == [("HDFC Bank", "Home Loan", "8.50% - 9.40%")]
    old = policy_facts.current(["HDFC Bank"], ["roi"], today=date(2026, 3, 1))
    assert [r["value"] for r in old] == ["9.10%"]                                # the rate in force then


def test_a_question_gets_the_approved_values_first(monkeypatch):
    policy_facts.import_csv(CSV)
    policy_facts.approve([f["fact_id"] for f in policy_facts.listing()])
    monkeypatch.setattr(retrieval, "vector_hits", lambda *a, **k: [])
    monkeypatch.setattr(retrieval, "keyword_hits", lambda *a, **k: [])
    prep = chat.Prepared("kb", "What is the HDFC home loan ROI?", {}, {}, None, [], [], "", use_llm=False)
    assert not prep.not_found and prep.refs[0]["source"] == "Structured policy facts"
    assert "Rate of interest: 8.50% - 9.40%" in prep.context and "HDFC HL rate sheet Aug 2026.pdf" in prep.context
    assert "9.10%" not in prep.context                                           # expired
    assert prep.finish("HDFC home loan ROI is **8.50% - 9.40%** [S1]")["unsupported_figures"] == []
    assert chat.Prepared("kb", "Axis bank home loan ROI?", {}, {}, None, [], [], "", use_llm=False).facts is None


def test_admin_api():
    c = _admin()
    r = c.post("/api/admin/kb/facts/import", files={"file": ("facts.csv", CSV.encode(), "text/csv")})
    assert r.status_code == 200 and r.json()["added"] == 3
    pending = c.get("/api/admin/kb/facts?status=pending").json()["facts"]
    assert c.post("/api/admin/kb/facts/approve", json={"ids": [f["fact_id"] for f in pending]}).json() == {"approved": 3}
    bad = c.post("/api/admin/kb/facts/import", files={"file": ("x.csv", b"lender,value\nA,1\n", "text/csv")})
    assert bad.status_code == 400 and "fact" in bad.json()["detail"]


def test_payout_grid_becomes_pending_facts(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [["Bank Name", "HOME LOAN (HL)", "LOAN AGAINST PROPERTY (LAP)", "SME"],
              ["Aadhar housing Finance", 0.01, 0.012, "-"], ["ICICI BANK", 0.0065, "-", "Case to case"]]:
        ws.append(r)
    for c in ("B2", "C2", "B3"):
        ws[c].number_format = "0.00%"
    wb.save(tmp_path / "Slab Book Secured July 2026.xlsx")
    r = policy_facts.extract_payout_grid(tmp_path / "Slab Book Secured July 2026.xlsx")
    assert r == {"rows": 3, "added": 3, "skipped": 0}
    rows = {(f["lender"], f["product"], f["value"], f["effective_from"], f["status"]) for f in policy_facts.listing()}
    assert rows == {("Aadhar housing Finance", "Home Loan", "1%", "2026-07-01", "pending"),
                    ("Aadhar housing Finance", "LAP", "1.2%", "2026-07-01", "pending"),
                    ("ICICI Bank", "Home Loan", "0.65%", "2026-07-01", "pending")}


def test_payout_grid_without_a_lender_header_and_with_remarks(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [["", "BUSINESS LOAN (BL)", "PERSONAL LOAN (PL)", "T & C/ Remarks"],
              ["Aditya Birla Finance Limited", "4%", "Refer Sheet", "Payout 0.5% lower for top-up"]]:
        ws.append(r)
    wb.save(tmp_path / "Slab Book Unsecured July 2026.xlsx")
    assert policy_facts.extract_payout_grid(tmp_path / "Slab Book Unsecured July 2026.xlsx")["added"] == 1
    f = policy_facts.listing()[0]
    assert (f["lender"], f["product"], f["value"]) == ("Aditya Birla Capital", "Business Loan", "4%")
