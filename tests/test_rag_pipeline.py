"""Grounded RAG pipeline: structure-aware extraction, versioning, query understanding, metadata filtering,
evidence validation, citations, the figure check, traces, diagnosis and golden tests."""
import time
import uuid

import pytest
from fastapi.testclient import TestClient

import app.rag.chat as chat
import app.rag.retrieval as retrieval
import app.rag.vectorstore as vectorstore
from app.core.config import get_settings
from app.main import app
from app.rag import entities, grounding
from app.services import accounts, kb_admin, kb_versions, rag_debug
from app.services import knowledge_base as kb
from app.services.ingestion import extract, load_and_chunk
from tests.test_knowledge_base import FakeStore


@pytest.fixture
def kbenv(monkeypatch, tmp_path):
    root = tmp_path / "Urban Money"
    root.mkdir()
    store = FakeStore()
    monkeypatch.setattr(get_settings(), "urban_money_dir", root)
    monkeypatch.setattr(get_settings(), "kb_index_db", tmp_path / "kb.db")
    monkeypatch.setattr(get_settings(), "kb_cache_dir", tmp_path / "cache")
    monkeypatch.setattr(vectorstore, "is_pinecone", lambda: True)
    monkeypatch.setattr(vectorstore, "get_vectorstore", lambda: store)
    monkeypatch.setattr(retrieval, "vector_hits", lambda q, k, lt, banks=None: [])     # keyword + exact search only
    kb_versions._invalidate()
    return root, store


POLICIES = {
    "HDFC Business Loan Policy.md": "# Loan Amount\nMaximum loan amount is ₹50 lakh for business loans.\n\n"
                                    "# Documents\nPAN, Aadhaar and 12 months bank statement.",
    "ICICI Business Loan Policy.md": "# Loan Amount\nMaximum loan amount is ₹75 lakh for business loans.",
    "HDFC Processing Notes.txt": "HDFC Bank business loan: required documents are PAN and GST returns.",
}


def _write(root, files):
    for name, text in files.items():
        (root / name).write_text(text, encoding="utf-8")


class FakeLLM:
    def __init__(self, answer="", rewrite=None):
        self.answer, self.rewrite, self.prompts = answer, rewrite, []

    def _text(self, prompt):
        return prompt if isinstance(prompt, str) else "\n".join(m.content for m in prompt)

    def stream(self, prompt):
        self.prompts.append(self._text(prompt))
        for piece in self.answer.split(" "):
            yield type("C", (), {"content": piece + " "})()

    def invoke(self, prompt):
        self.prompts.append(self._text(prompt))
        return type("R", (), {"content": self.answer})()


def _admin():
    c = TestClient(app)
    r = c.post("/api/auth/register", json={"email": f"rag-{uuid.uuid4().hex[:8]}@example.com", "name": "RAG Admin",
                                           "password": "correct-horse"})
    accounts.update_user(r.json()["user"]["id"], -1, role="admin")
    return c


# ------------------------------------------------------------------ extraction and chunking
def test_docx_keeps_headings_lists_and_tables_in_order(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("HDFC Business Loan Policy", 1)
    d.add_heading("Loan Amount", 2)
    d.add_paragraph("Maximum loan amount is ₹50 lakh.")
    d.add_paragraph("Self-employed only", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Tenure", "ROI", "5 years", "14%"
    d.add_heading("Eligibility", 2)
    d.add_paragraph("Minimum CIBIL 700.")
    f = tmp_path / "policy.docx"
    d.save(f)
    parts = extract(f)
    assert parts[0].heading == "HDFC Business Loan Policy > Loan Amount"
    assert parts[0].text == "Maximum loan amount is ₹50 lakh.\n- Self-employed only"
    assert parts[1].rows == ["Tenure: 5 years | ROI: 14%"] and parts[1].heading.endswith("Loan Amount")
    assert parts[2].heading == "HDFC Business Loan Policy > Eligibility" and "CIBIL 700" in parts[2].text


def test_pdf_keeps_pages_headings_and_drops_running_footer(tmp_path):
    import pymupdf
    doc = pymupdf.open()
    for i, (head, body) in enumerate([("Loan Amount", "Maximum loan amount is Rs 50 lakh."),
                                      ("Tenure", "Tenure up to 5 years."), ("Charges", "Processing fee 2%.")], 1):
        page = doc.new_page()
        page.insert_text((72, 100), head, fontsize=18)
        page.insert_text((72, 140), body, fontsize=10)
        page.insert_text((72, page.rect.height - 20), f"Confidential - page {i}", fontsize=8)
    f = tmp_path / "HDFC Business Loan Policy.pdf"
    doc.save(f)
    parts = extract(f)
    assert [(p.page, p.heading) for p in parts] == [(1, "Loan Amount"), (2, "Tenure"), (3, "Charges")]
    assert parts[0].text == "Maximum loan amount is Rs 50 lakh."
    assert not any("Confidential" in p.text for p in parts)                   # footer on every page dropped

    meta = {"title": "HDFC Business Loan Policy", "bank": "HDFC Bank", "product": "Business Loan", "version": 3}
    chunk = load_and_chunk(f, parts=parts, extra_meta=meta)[0]
    assert chunk.page_content.splitlines()[:3] == [
        "Source: HDFC Business Loan Policy.pdf | Page 1",
        "Document: HDFC Business Loan Policy | Bank: HDFC Bank | Product: Business Loan | Version: v3",
        "Section: Loan Amount"]
    assert chunk.metadata["page"] == 1 and chunk.metadata["section"] == "Loan Amount" and chunk.metadata["chunk_index"] == 0


def test_metadata_is_inferred_from_the_file_name():
    m = kb_versions.infer("Policies/HDFC Business Loan Policy.pdf")
    assert m == {"title": "HDFC Business Loan Policy", "bank": "HDFC Bank", "product": "Business Loan", "category": "Policy"}
    multi = kb_versions.infer("Code Book-5.xlsx", "Axis Bank 6159MUM | HDFC Bank 7788 | ICICI 9911")
    assert multi["bank"] == "" and multi["category"] == "Codes"                 # a multi-bank sheet has no bank


# ------------------------------------------------------------------ versioning
def test_versions_only_active_is_indexed_and_rollback_works(kbenv):
    root, store = kbenv
    c = _admin()
    up = lambda text: c.post("/api/admin/kb/upload", files=[("files", ("HDFC BL Policy.txt", text.encode()))])
    up("HDFC business loan maximum loan amount Rs 50 lakh.")
    _wait()
    up("HDFC business loan maximum loan amount Rs 60 lakh.")
    _wait()
    item = c.get("/api/admin/kb").json()["pending"][0]
    c.post(f"/api/admin/kb/pending/{item['id']}/decide", json={"action": "version", "doc_id": item["candidates"][0]["doc_id"]})
    _wait()
    texts = " ".join(d.page_content for d in store.docs.values())
    assert "60 lakh" in texts and "50 lakh" not in texts                       # the archived version is not indexed
    info = c.get("/api/admin/kb/document", params={"path": "HDFC BL Policy.txt"}).json()
    assert [(v["version"], v["status"], v["restorable"]) for v in info["versions"]] == [(2, "ACTIVE", True), (1, "ARCHIVED", True)]
    assert info["document"]["bank"] == "HDFC Bank"
    assert all(d.metadata["version"] == 2 and d.metadata["status"] == "ACTIVE" for d in store.docs.values())

    old = c.get("/api/admin/kb/versions/preview", params={"path": "HDFC BL Policy.txt", "version": 1}).json()
    assert "50 lakh" in old["chunks"][0] and old["status"] == "ARCHIVED"

    assert c.post("/api/admin/kb/versions/activate", json={"path": "HDFC BL Policy.txt", "version": 1}).json()["activated"] == 1
    _wait()
    texts = " ".join(d.page_content for d in store.docs.values())
    assert "50 lakh" in texts and "60 lakh" not in texts
    statuses = {v["version"]: v["status"] for v in c.get("/api/admin/kb/document", params={"path": "HDFC BL Policy.txt"}).json()["versions"]}
    assert statuses == {1: "ACTIVE", 2: "ARCHIVED"}

    c.post("/api/admin/kb/archive", json={"paths": ["HDFC BL Policy.txt"]})
    _wait()
    assert kb_versions.versions("HDFC BL Policy.txt")[-1]["status"] == "DISABLED"
    assert "HDFC BL Policy.txt" not in kb_versions.active_versions()


def test_metadata_edit_reindexes_with_new_tags(kbenv):
    root, store = kbenv
    _write(root, {"notes.txt": "Business loan maximum loan amount Rs 40 lakh."})
    c = _admin()
    c.post("/api/admin/kb/sync", json={})
    _wait()
    r = c.patch("/api/admin/kb/document", json={"path": "notes.txt", "bank": "Axis Bank", "product": "Business Loan"})
    assert r.status_code == 200 and r.json()["document"]["meta_source"] == "admin"
    _wait()
    d = next(iter(store.docs.values()))
    assert d.metadata["bank"] == "Axis Bank" and "Bank: Axis Bank" in d.page_content


def test_stale_chunk_of_an_old_version_is_never_retrieved(kbenv, monkeypatch):
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    from langchain_core.documents import Document
    stale = Document(id="stale-1", page_content="HDFC business loan maximum loan amount ₹90 lakh",
                     metadata={"corpus": "urban_money", "source": "HDFC Business Loan Policy.md", "sha": "000000000000",
                               "status": "ACTIVE", "bank": "HDFC Bank"})
    monkeypatch.setattr(retrieval, "vector_hits", lambda q, k, lt, banks=None: [(stale, 0.99)])
    r = retrieval.retrieve("What is the maximum HDFC business loan amount?", use_llm=False)
    assert "stale-1" not in [c.key for c in r.selected]
    assert next(c for c in r.candidates if c.key == "stale-1").dropped == "not the active version"


def _wait():
    for _ in range(200):
        if not kb_admin.busy():
            return
        time.sleep(0.05)
    raise AssertionError("knowledge-base job did not finish")


# ------------------------------------------------------------------ query understanding
def test_entities_and_metrics():
    assert entities.find_banks("pnb housing vs pnb") == ["PNB Housing Finance", "Punjab National Bank"]
    assert entities.find_products("Axis HL and LAP") == ["Home Loan", "LAP"]
    assert entities.find_metrics("What is the maximum HDFC business loan amount?") == ["max_amount"]
    assert entities.find_metrics("maximum tenure for a home loan") == ["tenure"]


def test_followup_is_resolved_from_the_conversation():
    history = [{"role": "user", "text": "What is the maximum loan amount for a business loan?"},
               {"role": "assistant", "text": "₹50 lakh."}]
    p = retrieval.understand("What about HDFC?", history, use_llm=False)
    assert p.followup and p.rewrite == "rules"
    assert p.banks == ["HDFC Bank"] and p.products == ["Business Loan"] and p.metrics == ["max_amount"]
    assert any("HDFC" in q and "maximum loan amount" in q for q in p.queries)


def test_precise_question_is_not_rewritten(monkeypatch):
    llm = FakeLLM()
    monkeypatch.setattr("app.rag.llm.get_llm", lambda kind="fast": llm)
    p = retrieval.understand("How much can I get from HDFC for a business loan?")
    assert p.rewrite == "none" and llm.prompts == [] and p.queries[0] == "How much can I get from HDFC for a business loan?"


def test_llm_rewrite_that_changes_the_meaning_is_rejected(monkeypatch):
    class Rewriter:
        def with_structured_output(self, schema, method=None):
            return self

        def with_retry(self, **kw):
            return self

        def invoke(self, prompt):
            return retrieval.QueryRewrite(standalone_question="What is the ICICI business loan limit?",
                                          search_query="ICICI business loan")

    monkeypatch.setattr("app.rag.llm.get_llm", lambda kind="fast": Rewriter())
    history = [{"role": "user", "text": "Tell me about HDFC business loans"}]
    p = retrieval.understand("and what is the limit there?", history)
    assert p.rewrite == "llm rejected: added ICICI Bank" and p.banks == ["HDFC Bank"]


# ------------------------------------------------------------------ filtering, reranking, validation
def test_bank_filter_and_evidence_check(kbenv):
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    r = retrieval.retrieve("What is the maximum HDFC business loan amount?", use_llm=False)
    assert r.validation["sufficient"] and len(r.passes) == 1
    assert r.selected[0].doc.metadata["source"] == "HDFC Business Loan Policy.md"
    assert r.selected[0].doc.metadata["section"] == "Loan Amount"
    icici = [c for c in r.candidates if c.doc.metadata["source"].startswith("ICICI")]
    assert icici and all(c.dropped == "different bank (ICICI Bank)" for c in icici)
    t = r.trace()
    assert t["plan"]["banks"] == ["HDFC Bank"] and t["passes"][0]["reranker"] == "lexical"
    assert all(k in t["candidates"][0] for k in ("vector_score", "keyword_score", "rerank", "score", "dropped"))


def test_missing_fact_expands_once_then_reports_not_found(kbenv, monkeypatch):
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    monkeypatch.setattr(retrieval, "vector_hits", lambda q, k, lt, banks=None: [])     # works, finds nothing
    r = retrieval.retrieve("What is the processing fee for an HDFC business loan?", use_llm=False)
    assert not r.validation["sufficient"] and [p["pass"] for p in r.passes] == [1, 2]
    assert "processing fee" in r.validation["reason"] and r.degraded == []


def test_a_failed_vector_search_is_retried_and_reported_not_as_missing(kbenv, monkeypatch):
    from app.rag import vectorstore
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    scales = []

    def timing_out(q, k, lt, banks=None):
        scales.append(vectorstore.timeout_scale.get())
        vectorstore.note_issue("timeout")
        return []
    monkeypatch.setattr(retrieval, "vector_hits", timing_out)
    r = retrieval.retrieve("What is the processing fee for an HDFC business loan?", use_llm=False)
    assert [p["pass"] for p in r.passes] == [1, "recovery", 2]
    assert retrieval.RECOVERY_TIMEOUT_SCALE in scales             # the recovery search waits longer
    assert r.degraded and r.trace()["degraded"] == r.degraded     # the outage is visible downstream
    assert vectorstore.issues.get() is None                       # nothing leaks into the next request


# ------------------------------------------------------------------ grounding
def test_figure_check_catches_changed_values():
    ctx = "Max loan amount: Rs. 50,00,000 | ROI 10.5% | Tenure 5 years | Code 6159MUM\nMax Loan (in Lakhs): 75"
    assert grounding.verify_figures("Up to **₹50 lakh** [S1] at 10.5% for 5 years, code 6159MUM.", ctx) == []
    assert grounding.verify_figures("Max ₹75 lakh [S1]", ctx) == []
    assert grounding.verify_figures("Up to ₹5 crore at 11% for 7 years, code 6159MUX.", ctx) == \
        ["₹5 crore", "11%", "7", "6159MUX"]
    assert grounding.verify_figures(grounding.NOT_FOUND, ctx) == []


def test_citations_map_to_the_cited_passages_only():
    refs = [{"ref": "S1", "title": "A"}, {"ref": "S2", "title": "B"}, {"ref": "S3", "title": "C"}]
    assert [r["ref"] for r in grounding.cited_refs("x [S3] y [S1, S3]", refs)] == ["S1", "S3"]
    assert all(r["uncited"] for r in grounding.cited_refs("no citations", refs))
    assert grounding.cited_refs(grounding.NOT_FOUND, refs) == []
    assert grounding.invalid_citations("see [S9]", refs) == ["S9"]


def test_stream_cites_sources_and_flags_unsupported_figures(kbenv, monkeypatch):
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    llm = FakeLLM("The maximum HDFC business loan amount is **₹5 crore** [S1].")
    monkeypatch.setattr(chat, "get_llm", lambda kind="fast": llm)
    c = _admin()
    from tests.test_chat_stream import _events
    ev = _events(c.post("/api/chat/stream", json={"mode": "kb", "question": "What is the maximum HDFC business loan amount?"}).text)
    sources = dict(ev)["sources"]
    assert [s["ref"] for s in sources] == ["S1"] and sources[0]["title"] == "HDFC Business Loan Policy"
    assert sources[0]["location"] == "Section: Loan Amount" and "₹50 lakh" in sources[0]["excerpt"]
    assert dict(ev)["verification"] == {"unsupported": ["₹5 crore"]}
    assert "document-grounded" in llm.prompts[0] and "[S1] HDFC Business Loan Policy" in llm.prompts[0]

    t = rag_debug.traces(limit=1)[0]
    assert t["question"] == "What is the maximum HDFC business loan amount?" and t["flagged"] == 1
    detail = c.get(f"/api/admin/rag/traces/{t['id']}", params={"expected_document": "HDFC Business Loan Policy",
                                                               "expected_answer": "₹50 lakh"}).json()
    d = detail["data"]
    assert d["generation"]["temperature"] == 0.0 and d["prompt"]["version"] == grounding.PROMPT_VERSION
    assert d["config"]["reranker"] == "lexical" and d["unsupported_figures"] == ["₹5 crore"]
    assert [s["status"] for s in detail["diagnosis"]][-1] == "problem"            # the LLM changed the figure
    assert detail["diagnosis"][-1]["stage"].startswith("7-8")


# ------------------------------------------------------------------ admin test bench and golden tests
def test_admin_test_bench_without_llm_evaluates_retrieval(kbenv):
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    c = _admin()
    r = c.post("/api/admin/rag/test", json={"question": "What is the maximum HDFC business loan amount?",
                                            "expected_document": "HDFC Business Loan Policy"}).json()
    assert r["llm_skipped"] and r["evaluation"]["retrieval"] == "GOOD" and r["evaluation"]["rank"] == 1
    assert [s["stage"] for s in r["diagnosis"]][:5] == ["1. Extraction", "2. Chunking", "3. Retrieval",
                                                       "4. Metadata filtering", "5. Reranking"]

    nf = c.post("/api/admin/rag/test", json={"question": "What is the processing fee for an HDFC business loan?"}).json()
    assert nf["not_found"] and nf["answer"] == chat.NOTHING_FOUND


def test_golden_tests_run_and_compare(kbenv, monkeypatch):
    root, _ = kbenv
    _write(root, POLICIES)
    kb.sync()
    c = _admin()
    assert TestClient(app).get("/api/admin/rag/golden").status_code == 401
    good = c.post("/api/admin/rag/golden", json={"question": "What is the maximum HDFC business loan amount?",
                                                 "expected_answer": "₹50 lakh", "expected_document": "HDFC Business Loan Policy"}).json()
    c.post("/api/admin/rag/golden", json={"question": "What is the maximum ICICI business loan amount?",
                                          "expected_answer": "₹75 lakh", "expected_document": "ICICI Business Loan Policy"})
    monkeypatch.setattr(chat, "get_llm", lambda kind="fast": FakeLLM("It is ₹50 lakh [S1]."))
    rid = rag_debug.run_golden("manual")
    run = c.get(f"/api/admin/rag/golden/runs/{rid}").json()
    assert run["total"] == 2 and run["passed"] == 1 and run["config"]["reranker"] == "lexical"
    by_q = {x["question"]: x for x in run["results"]}
    assert by_q["What is the maximum HDFC business loan amount?"]["passed"]
    icici = by_q["What is the maximum ICICI business loan amount?"]
    assert icici["retrieval"] == "GOOD" and icici["answer"] == "WRONG" and not icici["passed"]

    c.put(f"/api/admin/rag/golden/{good['id']}", json={"question": good["question"], "expected_answer": "₹60 lakh",
                                                       "expected_document": "HDFC Business Loan Policy"})
    run2 = c.get(f"/api/admin/rag/golden/runs/{rag_debug.run_golden('manual')}").json()
    assert run2["passed"] == 0 and any(x["regression"] for x in run2["results"])
    assert c.get("/api/admin/rag/config").json()["hash"] == run2["config_hash"]
    assert c.delete(f"/api/admin/rag/golden/{good['id']}").status_code == 200


def test_answer_matching():
    assert rag_debug.answer_matches("₹50 lakh", "The limit is Rs. 50,00,000 [S1].")
    assert not rag_debug.answer_matches("₹50 lakh", "The limit is ₹5 crore.")
    assert rag_debug.answer_matches("Asha Patil", "Contact **Asha Patil** on 9000000000")
